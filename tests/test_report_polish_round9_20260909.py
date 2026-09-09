from datetime import date
from pathlib import Path

from docx import Document
from PIL import Image

from apps.reports import exporting
from apps.reports.forms import BOOLEAN_REPORT_FIELDS, PERSISTED_REPORT_FIELDS, ReportCreateForm
from apps.reports.narratives import section_enabled


def test_report_builder_contains_requested_controls_and_styles():
    root = Path(__file__).resolve().parents[1]
    report_template = (root / "templates/reports/report_list.html").read_text()
    version_template = (root / "templates/reports/version_detail.html").read_text()
    project_template = (root / "templates/reports/project_list.html").read_text()
    styles = (root / "static/reports/app.css").read_text()

    assert "metrika_goals_humans_only" in report_template
    assert "Роботность: только люди" in report_template
    assert "nested-checkbox-option" in report_template
    assert "search-engine-details" in report_template
    assert 'name="is_draft"' not in version_template
    assert 'class="compact-add"' in project_template
    assert '<svg viewBox="0 0 24 24"' in project_template
    assert ".compact-add svg" in styles
    assert "grid-template-columns:repeat(4,max-content)" in styles


def test_goal_robotness_defaults_to_humans_and_is_persisted():
    form = ReportCreateForm(project=None)

    assert form.fields["metrika_goals_humans_only"].initial is True
    assert "metrika_goals_humans_only" in PERSISTED_REPORT_FIELDS
    assert "metrika_goals_humans_only" in BOOLEAN_REPORT_FIELDS


def test_metrika_detail_variants_follow_general_robotness_exactly():
    payload = {
        "display_options": {"metrika_robotness": "all", "metrika_search_segment": True},
        "calculated": {
            "sources": {
                "sources": {
                    "yandex_metrika": {
                        "period_details": [
                            {
                                "period_start": "2026-08-01",
                                "period_end": "2026-08-31",
                                "payload": {
                                    "detail_variants": {
                                        "search": {
                                            "humans": {"search_engines": [{"visits": "10"}]},
                                            "all": {"search_engines": [{"visits": "14"}]},
                                        }
                                    },
                                    "traffic_source_variants": {
                                        "humans": {
                                            "rows": [{"code": "search", "visits": "10"}],
                                            "total": {"visits": "10"},
                                        },
                                        "all": {
                                            "rows": [{"code": "search", "visits": "14"}],
                                            "total": {"visits": "14"},
                                        },
                                    },
                                },
                            }
                        ]
                    }
                }
            }
        },
    }

    assert exporting._metrika_period_rows(payload, "search_engines")[0]["rows"][0]["visits"] == "14"
    assert (
        exporting._metrika_period_rows(payload, "traffic_source_details")[0]["total"]["visits"]
        == "14"
    )
    payload["display_options"]["metrika_robotness"] = "humans"
    assert exporting._metrika_period_rows(payload, "search_engines")[0]["rows"][0]["visits"] == "10"


def test_all_robotness_never_falls_back_to_legacy_human_details():
    payload = {
        "display_options": {"metrika_robotness": "all", "metrika_search_segment": True},
        "calculated": {
            "sources": {
                "sources": {
                    "yandex_metrika": {
                        "period_details": [
                            {
                                "period_start": "2026-08-01",
                                "period_end": "2026-08-31",
                                "payload": {
                                    "search_engines": [{"visits": "10"}],
                                    "traffic_source_details": [{"code": "search", "visits": "10"}],
                                },
                            }
                        ]
                    }
                }
            }
        },
    }

    assert exporting._metrika_period_rows(payload, "search_engines")[0]["rows"] == []
    assert exporting._metrika_period_rows(payload, "traffic_source_details")[0]["rows"] == []


def test_goal_robotness_is_independent_from_general_robotness(monkeypatch):
    captured = []
    monkeypatch.setattr(exporting, "_goal_card", lambda _doc, goal, _periods: captured.append(goal))
    payload = {
        "display_options": {
            "configuration_version": 3,
            "include_metrika": True,
            "include_metrika_search_engines": False,
            "include_metrika_geography": False,
            "include_metrika_landing_pages": False,
            "include_metrika_landing_page_comparison": False,
            "include_metrika_url_groups": False,
            "include_metrika_sections": False,
            "include_metrika_categories": False,
            "include_metrika_goals": True,
            "metrika_robotness": "humans",
            "metrika_goals_humans_only": False,
            "metrika_goals_quarter": False,
            "metrika_search_segment": True,
        },
        "calculated": {
            "sources": {
                "sources": {
                    "yandex_metrika": {
                        "period_details": [
                            {
                                "period_start": "2026-08-01",
                                "period_end": "2026-08-31",
                                "payload": {
                                    "goals_by_segment": {
                                        "search": {
                                            "humans": [{"goal_id": "1", "visits": "10"}],
                                            "all": [{"goal_id": "1", "visits": "14"}],
                                        }
                                    }
                                },
                            }
                        ]
                    }
                }
            }
        },
    }

    document = Document()
    exporting._configure_document(document, "site.test", date(2026, 8, 1))
    exporting._render_metrika(document, payload, {})

    assert captured[0]["visits"] == "14"


def test_comparison_periods_and_table_are_previous_then_report(monkeypatch):
    captured = {}
    monkeypatch.setattr(
        exporting,
        "_period_pills_image",
        lambda items: captured.setdefault("items", items),
    )
    monkeypatch.setattr(exporting, "_add_report_picture", lambda *_args, **_kwargs: True)
    periods = [
        {"period_start": "2026-07-01", "period_end": "2026-07-31"},
        {"period_start": "2026-08-01", "period_end": "2026-08-31"},
    ]

    exporting._comparison_period_pills(Document(), periods)

    assert "июля" in captured["items"][0][0]
    assert captured["items"][0][1] == "A"
    assert "августа" in captured["items"][2][0]
    assert captured["items"][2][1] == "B"

    table_document = Document()
    exporting._configure_document(table_document, "site.test", date(2026, 8, 1))
    table = exporting._metrika_detail_table(
        table_document,
        [("Google", {"visits": 120}, {"visits": 100})],
        first_header="Поисковая система",
        metrics=("visits",),
        include_total=False,
    )
    assert [cell.text.splitlines()[0] for cell in table.rows[1].cells] == [
        "Google",
        "100",
        "120",
    ]
    assert len(table.rows[1].cells[1].paragraphs) == 1
    assert table.rows[1].cells[2].paragraphs[1].text == "20,00%"


def test_zero_metrika_series_are_omitted_from_chart_and_legend(monkeypatch):
    monkeypatch.setattr(exporting, "_save_figure", lambda figure: figure)
    figure = exporting._metrika_sources_chart(
        {
            "search": {
                "series": [
                    {"month": "2026-07-01", "value": 100},
                    {"month": "2026-08-01", "value": 110},
                ],
                "change": {"current": 110},
            },
            "advertising": {
                "series": [
                    {"month": "2026-07-01", "value": 0},
                    {"month": "2026-08-01", "value": 0},
                ],
                "change": {"current": 0},
            },
        }
    )

    legend_labels = [text.get_text() for text in figure.axes[0].get_legend().get_texts()]
    assert len(figure.axes[0].lines) == 1
    assert all("рекламе" not in label for label in legend_labels)


def test_generated_images_are_print_quality_and_docx_disables_compression():
    image = exporting._period_pills_image(
        [("1—30 июня ⌄", "A"), ("⇄", "swap"), ("1—31 июля ⌄", "B")]
    )
    with Image.open(image) as rendered:
        assert rendered.info.get("dpi", (0, 0))[0] >= 449
        assert rendered.width >= 2500

    document = Document()
    exporting._configure_document(document, "site.test", date(2026, 8, 1))
    settings_xml = document.settings._element.xml
    assert "doNotCompressPictures" in settings_xml
    assert "defaultImageDpi" in settings_xml
    assert 'w14:val="450"' in settings_xml


def test_goal_icon_is_larger_and_rendered_at_high_resolution(monkeypatch):
    monkeypatch.setattr(exporting, "_save_figure", lambda figure: figure)
    figure = exporting._metrika_goal_image(
        {"goal_id": "1", "name": "Телефон", "type": "call"},
        [{"period_start": "2026-08-01", "conversion_rate": 1, "visits": 2, "reaches": 3}],
    )

    icon_bounds = figure.axes[0].get_position().bounds
    assert icon_bounds[2] >= 0.02
    assert icon_bounds[3] >= 0.069


def test_landing_comparison_is_disabled_by_its_main_checkbox():
    payload = {
        "display_options": {
            "configuration_version": "3",
            "include_metrika": True,
            "include_metrika_landing_page_comparison": False,
            "metrika_url_segments": {
                "landing_comparison_subsections": [
                    {"name": "УЗИ", "patterns": ["https://site.test/uzi/*"]}
                ]
            },
        }
    }

    assert section_enabled(payload, "metrika_landing_page_comparison") is False
