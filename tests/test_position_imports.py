from datetime import date
from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook

from apps.imports.models import FileImportSegment, ImportBatch
from apps.imports.parser import parse_position_file, parse_position_history_xlsx
from apps.metrics.models import KeywordPosition, RankingSnapshot
from apps.projects.models import Project

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def isolated_media_root(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path / "media"


@pytest.fixture
def staff_client(client, django_user_model):
    user = django_user_model.objects.create_superuser(
        username="admin", email="admin@example.com", password="safe-test-password"
    )
    client.force_login(user)
    return client


@pytest.fixture
def project():
    return Project.objects.create(name="Demo", domain="example.com")


def csv_upload(content, name="positions.csv"):
    return SimpleUploadedFile(name, content.encode("utf-8"), content_type="text/csv")


def upload_positions(client, project, source_file):
    return client.post(
        reverse("imports:upload"),
        {
            "project": str(project.id),
            "snapshot_date": "2026-07-31",
            "search_engine": "yandex",
            "region": "Москва",
            "ranking_depth": "20",
            "source_file": source_file,
        },
    )


def history_xlsx_upload(name="history.xlsx", *, later_position=3):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Позиции"
    sheet.append(["Ключевое слово", "Страница", "Частота", "17.08.2026", "06.09.2026"])
    sheet.append(["Группа: плитка", "", "", "", ""])
    sheet.append(["низкая частота", "https://example.com/low", 10, 8, later_position])
    sheet.append(["высокая частота", "https://example.com/high", 500, 4, "-"])
    stream = BytesIO()
    workbook.save(stream)
    return SimpleUploadedFile(
        name,
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def test_history_xlsx_parses_groups_dates_missing_positions_and_frequency():
    upload = history_xlsx_upload()
    preview = parse_position_history_xlsx(upload.name, upload.read())

    assert preview.dates == (date(2026, 8, 17), date(2026, 9, 6))
    assert preview.keyword_count == 2
    september = preview.rows_by_date[date(2026, 9, 6)]
    assert [row["group_name"] for row in september] == ["плитка", "плитка"]
    assert september[0]["position_value"] == 3
    assert september[1]["position_status"] == KeywordPosition.Status.NOT_FOUND


def test_file_import_creates_independent_idempotent_segments_and_standard_calendars(
    staff_client, project
):
    url = reverse("imports:project-settings", args=[project.id])
    first = staff_client.post(
        url,
        {
            "search_engine": "yandex",
            "region": "Москва",
            "calculate_visibility": "on",
            "source_file": history_xlsx_upload(),
        },
    )
    assert first.status_code == 302
    segment = FileImportSegment.objects.get(project=project, search_engine="yandex")
    assert (
        RankingSnapshot.objects.filter(
            project=project, topvisor_configuration_id=segment.configuration_id
        ).count()
        == 2
    )
    september = RankingSnapshot.objects.get(
        project=project,
        topvisor_configuration_id=segment.configuration_id,
        snapshot_date=date(2026, 9, 6),
    )
    assert september.depth_source == RankingSnapshot.DepthSource.FILE_IMPORT
    assert september.visibility is not None

    staff_client.post(
        url + f"?edit={segment.id}",
        {
            "search_engine": "yandex",
            "region": "Москва",
            "calculate_visibility": "on",
            "source_file": history_xlsx_upload(later_position=1),
        },
    )
    assert (
        RankingSnapshot.objects.filter(
            project=project, topvisor_configuration_id=segment.configuration_id
        ).count()
        == 2
    )
    september.refresh_from_db()
    assert september.positions.get(normalized_query="низкая частота").position_value == 1

    staff_client.post(
        url,
        {
            "search_engine": "google",
            "region": "Москва",
            "source_file": history_xlsx_upload("google.xlsx"),
        },
    )
    assert FileImportSegment.objects.filter(project=project).count() == 2
    response = staff_client.get(reverse("reports:report-list", args=[project.id]))
    assert response.status_code == 200
    assert len(response.context["form"].configuration_date_fields) == 2
    assert response.context["form"].initial["include_visibility_table"] is True
    assert 'class="calendar-pair file-import-calendars calendar-count-2"' in (
        response.content.decode()
    )


def test_csv_preview_and_confirmation_create_normalized_positions(staff_client, project):
    source = (
        "Запрос;Позиция;Частотность;Группа;URL\n"
        "Купить камень;3;120;Каталог;https://WWW.Example.com/catalog/#section\n"
        "Доставка;>100;0;;/delivery/?from=seo#details\n"
    )
    response = upload_positions(staff_client, project, csv_upload(source))
    assert response.status_code == 302

    batch = ImportBatch.objects.get()
    assert batch.total_rows == 2
    assert batch.valid_rows == 2
    assert batch.error_rows == 0
    assert batch.status == ImportBatch.Status.PREVIEW
    assert batch.source_file.storage.exists(batch.source_file.name)

    response = staff_client.post(reverse("imports:confirm", args=[batch.id]))
    assert response.status_code == 302
    batch.refresh_from_db()
    assert batch.status == ImportBatch.Status.IMPORTED

    snapshot = RankingSnapshot.objects.get(import_batch=batch)
    assert snapshot.tracked_keyword_count == 2
    assert snapshot.ranking_depth == 20
    ranked = KeywordPosition.objects.get(normalized_query="купить камень")
    assert ranked.frequency == 120
    assert ranked.position_value == 3
    assert ranked.position_status == KeywordPosition.Status.RANKED
    assert ranked.normalized_target_url == "https://example.com/catalog/"
    beyond = KeywordPosition.objects.get(normalized_query="доставка")
    assert beyond.frequency == 1
    assert beyond.position_status == KeywordPosition.Status.BEYOND_100
    assert beyond.normalized_target_url == "/delivery/?from=seo"


def test_frequency_column_is_required(staff_client, project):
    source = "Запрос;Позиция\nКупить камень;3\n"
    response = upload_positions(staff_client, project, csv_upload(source))
    assert response.status_code == 200
    assert "Частотность" in response.content.decode("utf-8")
    assert not ImportBatch.objects.exists()


def test_missing_frequency_creates_row_error_and_blocks_confirmation(staff_client, project):
    source = "Запрос;Позиция;Частотность\nКупить камень;3;\n"
    response = upload_positions(staff_client, project, csv_upload(source))
    assert response.status_code == 302
    batch = ImportBatch.objects.get()
    assert batch.valid_rows == 0
    assert batch.error_rows == 1
    error = batch.row_errors.get()
    assert error.code == "invalid_frequency"

    staff_client.post(reverse("imports:confirm", args=[batch.id]))
    batch.refresh_from_db()
    assert batch.status == ImportBatch.Status.PREVIEW
    assert not RankingSnapshot.objects.exists()


def test_duplicate_file_with_same_parameters_is_idempotent(staff_client, project):
    source = "Запрос;Позиция;Частотность\nКупить камень;3;120\n"
    first = upload_positions(staff_client, project, csv_upload(source))
    second = upload_positions(staff_client, project, csv_upload(source))
    assert first.status_code == 302
    assert second.status_code == 302
    assert ImportBatch.objects.count() == 1


def test_duplicate_query_in_same_group_is_reported(staff_client, project):
    source = (
        "Запрос;Позиция;Частотность;Группа\n"
        "Купить камень;3;120;Каталог\n"
        "купить  КАМЕНЬ;4;100;каталог\n"
    )
    upload_positions(staff_client, project, csv_upload(source))
    batch = ImportBatch.objects.get()
    assert batch.valid_rows == 1
    assert batch.error_rows == 1
    assert batch.row_errors.get().code == "duplicate_query"


def test_xlsx_preview_is_supported(staff_client, project):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Служебная строка Topvisor"])
    worksheet.append(["Запрос", date(2026, 7, 31), "Частотность [WS]", "Группа"])
    worksheet.append(["Лестница из гранита", 7, 45, "Лестницы"])
    buffer = BytesIO()
    workbook.save(buffer)
    source_file = SimpleUploadedFile(
        "positions.xlsx",
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    response = upload_positions(staff_client, project, source_file)
    assert response.status_code == 302
    batch = ImportBatch.objects.get()
    assert batch.valid_rows == 1
    assert batch.error_rows == 0
    assert batch.preview_payload[0]["frequency"] == 45


@pytest.mark.parametrize("filename", ["positions.csv", "positions.xlsx"])
def test_zero_frequency_is_normalized_to_one_for_csv_and_xlsx(filename):
    if filename.endswith(".csv"):
        preview = parse_position_file(filename, "Запрос;Позиция;Частотность\nSEO;3;0\n".encode())
    else:
        workbook = Workbook()
        workbook.active.append(["Запрос", "Позиция", "Частотность"])
        workbook.active.append(["SEO", 3, 0])
        data = BytesIO()
        workbook.save(data)
        preview = parse_position_file(filename, data.getvalue())
    assert preview.errors == []
    assert preview.valid_rows[0]["frequency"] == 1


@pytest.mark.parametrize("frequency", ["", "-1", "invalid"])
def test_invalid_frequency_still_blocks_file_import(frequency):
    preview = parse_position_file(
        "positions.csv",
        f"Запрос;Позиция;Частотность\nSEO;3;{frequency}\n".encode(),
    )
    assert preview.valid_rows == []
    assert preview.errors[0].code == "invalid_frequency"


def test_more_than_3000_rows_is_supported():
    rows = ["Запрос;Позиция;Частотность"]
    rows.extend(f"Запрос {index};1;10" for index in range(3001))
    preview = parse_position_file("positions.csv", "\n".join(rows).encode("utf-8"))

    assert preview.total_rows == 3001
    assert len(preview.valid_rows) == 3001
    assert preview.errors == []


def test_import_pages_require_staff_authentication(client):
    response = client.get(reverse("imports:upload"))
    assert response.status_code == 302
    assert reverse("admin:login") in response.url


def test_staff_user_can_download_csv_template(staff_client):
    response = staff_client.get(reverse("imports:template"))
    assert response.status_code == 200
    assert response["Content-Type"] == "text/csv; charset=utf-8"
    assert response.content.decode("utf-8-sig") == "Запрос;Позиция;Частотность;Группа;URL\r\n"
